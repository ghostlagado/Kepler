from datetime import date, datetime

from openpyxl import load_workbook

from domain.hora_a_hora import build_report
from reports.hora_a_hora_xlsx import export


def test_export_matches_horaxhora_layout(tmp_path):
    report = build_report(
        "L062",
        date(2026, 10, 20),
        datetime(2026, 10, 20, 18, 0),
        values_by_code={"572-": 12.5, "280-": 7.4},
        metas_by_dept={"beleza": 2078.82, "basket": 363.29},
    )
    path = export(report, tmp_path / "saida.xlsx")
    ws = load_workbook(path).active
    assert ws["A2"].value == "HORA X HORA"
    assert ws["A3"].value == "DATA 20/10/2026"
    assert ws["A7"].value == "Beleza"
    assert ws["C7"].value == 2078.82
    assert ws["E7"].value == 7.4
    assert ws["G7"].value == "=E7-C7"
    assert ws["A16"].value == "Relógios"
    assert ws["G16"].value == "=E16-C16"
    total_row = 17
    assert ws[f"A{total_row}"].value == "TOTAL LOJA DIA"
    assert ws[f"C{total_row}"].value == "=SUM(C7:C16)"
    assert ws[f"E{total_row}"].value == "=SUM(E7:E16)"
