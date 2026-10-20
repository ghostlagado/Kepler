"""Exportação compatível com a planilha histórica horaxhora.xlsx do PDV_V2."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from domain.hora_a_hora import HourlyReport

_MONEY = "#,##0.00"
# Os dados começam na linha 7 porque a planilha histórica do PDV_V2 reserva
# as linhas 1-4 para cabeçalho e as 5-6 para folga — manter para que a
# exportação seja colada direto no acompanhamento antigo.
_FIRST_ROW = 7


def export(report: HourlyReport, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "HxH"

    ws["A1"] = "ACOMPANHAMENTO DE VENDAS"
    ws["A2"] = "HORA X HORA"
    ws["A3"] = f"DATA {report.report_date:%d/%m/%Y}"
    ws["A4"] = "DEPARTAMENTOS"
    ws["C4"] = "META DIA"
    ws["E4"] = "REALIZADO"
    ws["G4"] = "GAP"
    for cell in ("A4", "C4", "E4", "G4"):
        ws[cell].font = Font(bold=True)

    for offset, result in enumerate(report.results):
        row = _FIRST_ROW + offset
        ws[f"A{row}"] = result.department.label
        ws[f"C{row}"] = result.meta
        ws[f"E{row}"] = result.realizado
        ws[f"G{row}"] = f"=E{row}-C{row}"
        for col in ("C", "E", "G"):
            ws[f"{col}{row}"].number_format = _MONEY

    total_row = _FIRST_ROW + len(report.results)
    last_row = total_row - 1
    ws[f"A{total_row}"] = "TOTAL LOJA DIA"
    ws[f"A{total_row}"].font = Font(bold=True)
    ws[f"C{total_row}"] = f"=SUM(C{_FIRST_ROW}:C{last_row})"
    ws[f"E{total_row}"] = f"=SUM(E{_FIRST_ROW}:E{last_row})"
    ws[f"G{total_row}"] = f"=E{total_row}-C{total_row}"
    for col in ("C", "E", "G"):
        ws[f"{col}{total_row}"].number_format = _MONEY

    wb.save(path)
    wb.close()
    return path
