"""Leitura das metas diárias da planilha corporativa (layout do PDV_V2).

O arquivo tem três linhas de cabeçalho/filtro; os dados começam na linha 4
com colunas Data, Local, Grupo Operações e Meta. Linhas de outras datas ou
lojas são ignoradas, e grupos sem departamento conhecido são descartados.
"""

from datetime import date, datetime

from openpyxl import load_workbook

from domain.departments import DEPARTMENTS

_GROUP_TO_DEPARTMENT = {dept.meta_group: dept.key for dept in DEPARTMENTS}


class MetasError(Exception):
    """Planilha de metas ausente, ilegível ou em formato inesperado."""


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return None


def load_daily_metas(path, store: str, report_date: date) -> dict[str, float]:
    """Metas do dia por departamento (chave = Department.key)."""
    try:
        planilha = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise MetasError(f"não foi possível abrir {path}: {exc}") from exc
    try:
        ws = planilha.active
        metas: dict[str, float] = {}
        for row in ws.iter_rows(min_row=4, max_col=4, values_only=True):
            dia, loja, grupo, valor = row
            if _as_date(dia) != report_date or loja != store:
                continue
            dept_key = _GROUP_TO_DEPARTMENT.get(grupo)
            if dept_key is None or not isinstance(valor, (int, float)):
                continue
            metas[dept_key] = metas.get(dept_key, 0.0) + float(valor)
        return metas
    finally:
        planilha.close()
